"""Reading a .pptx for what it *means*, not for how it looks.

A deck exported to PNGs is a picture of a deck. It cuts into a video as a
still, and a still on screen for twenty seconds is where attention goes to die.
So this module does not treat the file as artwork. It treats it as the only
written record of what the presenter decided to say, in order, with emphasis —
and hands that record to a renderer that draws it again with motion.

What comes out is deliberately structural: every text run with its geometry,
its size relative to the rest of the slide, and a guess at its role. The guess
is a *guess*; `edit.plan` is free to overrule it. But the geometry is not, and
that is what makes the guess worth having — in every deck ever made, the
biggest text on the slide is the point of the slide, and the small line pinned
to the bottom is the aside.

PowerPoint itself, if installed, also exports reference PNGs. Nothing in the
pipeline renders those into the film. They exist so a human can put the
original next to the rebuild and see what was dropped.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

EMU = 914400.0  # english metric units per inch, the unit pptx stores geometry in


def _text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return shape.text_frame.text.replace("\x0b", "\n").strip()


def _largest_size(shape) -> float:
    """The biggest explicit font size on the shape, in points.

    Runs frequently inherit their size from the layout rather than stating it,
    which reads back as ``None``. Falling back to 18 keeps such a run in the
    ordering without letting it pretend to be a headline.
    """
    sizes: list[float] = []
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size is not None:
                sizes.append(run.font.size.pt)
    return max(sizes) if sizes else 18.0


def _bold(shape) -> bool:
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.bold:
                return True
    return False


def _colour(shape) -> str | None:
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            try:
                rgb = run.font.color.rgb
            except (AttributeError, TypeError, ValueError):
                continue
            if rgb is not None:
                return f"#{rgb}"
    return None


def _role(block: dict[str, Any], blocks: list[dict[str, Any]]) -> str:
    """A first guess, from geometry and typography alone.

    Ordered most-confident first. `kicker` before `headline` because a kicker
    ("STEP 3") is short, upper-case and pinned to the top, and would otherwise
    be mistaken for a headline on a slide whose real headline is set small.
    """
    text, size, top = block["text"], block["size"], block["top"]
    biggest = max(b["size"] for b in blocks)
    words = len(text.split())

    if top < 0.22 and words <= 4 and text.upper() == text and len(text) < 24:
        return "kicker"
    if size >= biggest * 0.92 and words <= 14:
        return "headline"
    if top > 0.78 and size <= biggest * 0.45:
        return "footnote"
    if re.fullmatch(r"\d{1,2}", text.strip()):
        return "index"
    if re.fullmatch(r"[>›→»]", text.strip()):
        return "arrow"
    if "\n" in text or words > 14:
        return "body"
    return "item"


def read(path: Path) -> dict[str, Any]:
    """The whole deck, as structure."""
    from pptx import Presentation

    deck = Presentation(str(path))
    width, height = deck.slide_width, deck.slide_height
    slides: list[dict[str, Any]] = []

    for index, slide in enumerate(deck.slides, start=1):
        blocks: list[dict[str, Any]] = []
        decorations = 0
        for shape in slide.shapes:
            text = _text(shape)
            if not text:
                # Bars, pills and connectors. Counted rather than described:
                # how *many* there are says whether the slide was a grid, which
                # is a useful hint. What colour they were is not.
                decorations += 1
                continue
            blocks.append({
                "text": text,
                "size": _largest_size(shape),
                "bold": _bold(shape),
                "colour": _colour(shape),
                "left": round((shape.left or 0) / width, 3),
                "top": round((shape.top or 0) / height, 3),
                "width": round((shape.width or 0) / width, 3),
                "height": round((shape.height or 0) / height, 3),
            })

        # Reading order: down the slide, then across. Shapes come out of the
        # file in creation order, which is the order somebody happened to draw
        # them in and bears no relation to how it is read.
        blocks.sort(key=lambda b: (round(b["top"], 2), b["left"]))
        for block in blocks:
            block["role"] = _role(block, blocks)

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        slides.append({
            "index": index,
            "blocks": blocks,
            "decorations": decorations,
            "notes": notes,
            "text": "\n".join(b["text"] for b in blocks),
        })

    return {
        "source": str(path),
        "aspect": f"{width}:{height}",
        "ratio": round(width / height, 4),
        "slides": slides,
    }


def outline(deck: dict[str, Any]) -> str:
    """The deck as something a person — or a model — can read in one screen."""
    lines: list[str] = []
    for slide in deck["slides"]:
        lines.append(f"## slide {slide['index']}")
        for block in slide["blocks"]:
            body = block["text"].replace("\n", " / ")
            lines.append(f"  [{block['role']:<8}] {body}")
        if slide["notes"]:
            lines.append(f"  (notes) {slide['notes']}")
    return "\n".join(lines)


def export_png(path: Path, outdir: Path) -> list[Path]:
    """Reference stills, via PowerPoint if it is here. Never used in the cut.

    Failure is not an error. The deck's *content* is what the pipeline needs
    and that came out of the XML; these are for a human comparing the rebuild
    against the original.
    """
    outdir.mkdir(parents=True, exist_ok=True)
    try:
        import comtypes.client
    except ImportError:
        print("  reference PNGs skipped (pip install comtypes)")
        return []
    try:
        app = comtypes.client.CreateObject("PowerPoint.Application")
        deck = app.Presentations.Open(str(path.resolve()), WithWindow=False)
        deck.SaveAs(str(outdir.resolve()), 18)  # ppSaveAsPNG
        deck.Close()
        app.Quit()
    except Exception as exc:  # COM raises a zoo of things, none of them useful
        print(f"  reference PNGs skipped ({exc.__class__.__name__})")
        return []
    found = sorted(outdir.glob("*.PNG")) + sorted(outdir.glob("*.png"))
    print(f"  reference PNGs         {len(found)}")
    return found


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="demoforge.deck", description=__doc__)
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--out", default=None, help="write deck.json here")
    ap.add_argument("--png", default=None, help="also export reference stills here")
    args = ap.parse_args(argv)

    deck = read(Path(args.pptx))
    print(outline(deck))
    if args.out:
        Path(args.out).write_text(json.dumps(deck, indent=1), encoding="utf-8")
    if args.png:
        export_png(Path(args.pptx), Path(args.png))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
